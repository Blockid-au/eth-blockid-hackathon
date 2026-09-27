"""Valuation v5 API (studio/projections.py, studio/finalise.py, guards in routes.py / offerings.py) against a
throwaway Postgres (TEST_DATABASE_URL; skipped when unset): flag off = 404 + unchanged behaviour; projections upload
-> confirm -> deterministic re-value; approval; finalise within +/-20 % (note), beyond (reason + a different platform
admin); low-confidence override; expiry / stale finals; company creation defaults and refusals; four-eyes assumption
changes; offering guards; /verify of a stored v5 report."""
import io
import json
from datetime import UTC, datetime, timedelta

import openpyxl
import pytest
from fastapi import HTTPException
from psycopg.types.json import Jsonb
from test_studio import ADMIN_KEY, USER, USER_KEY, needs_db, siwe_login, studio_env  # noqa: F401 (fixture)

from blockid_agents.studio import finalise as fin
from blockid_agents.studio import verify
from blockid_agents.studio.report_hash import canonical_json, report_hash_from_row, report_view_from_row
from blockid_agents.tools import projections as pj

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def v5_on(monkeypatch):
    monkeypatch.setenv("VALUATION_V5", "1")
    monkeypatch.delenv("VALUATION_V5_REQUIRE_FINAL", raising=False)


@pytest.fixture
def v5_off(monkeypatch):
    monkeypatch.setenv("VALUATION_V5", "0")


def forecast_xlsx(growth=0.4) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(pj.template_xlsx()))
    rev = [300e3, 480e3]
    for _ in range(5):
        rev.append(rev[-1] * (1 + growth))
    data = {"revenue": rev, "cogs": [r * 0.28 for r in rev], "opex": [r * 0.9 for r in rev[:3]] + [r * 0.6 for r in
                                                                                                   rev[3:]],
            "d_and_a": [r * 0.02 for r in rev], "capex": [r * 0.03 for r in rev]}
    for row in wb["Projections"].iter_rows(min_row=2):
        if row[0].value in data:
            for j, v in enumerate(data[row[0].value]):
                row[2 + j].value = v
    for row in wb["Company"].iter_rows(min_row=2):
        if row[0].value == "cash":
            row[2].value = 300000
        if row[0].value == "planned_raise":
            row[2].value = 1000000
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def multipart(data: bytes, name: str, field: str = "file") -> tuple[bytes, str]:
    b = "----blockidtest"
    body = (f"--{b}\r\nContent-Disposition: form-data; name=\"{field}\"; filename=\"{name}\"\r\n"
            f"Content-Type: {XLSX}\r\n\r\n").encode() + data + f"\r\n--{b}--\r\n".encode()
    return body, f"multipart/form-data; boundary={b}"


def run_valuation(env, client) -> str:
    vid = client.post("/v1/studio/valuations", json={"url": "agritrace.example"}).json()["id"]
    assert env["runner"].drain() == 1
    return vid


def second_admin(env):
    """The password admin (a different platform admin than the ADMIN wallet), password changed first."""
    p = env["client"]()
    assert p.post("/v1/auth/login", json={"username": "admin", "password": "admin"}).status_code == 200
    assert p.post("/v1/auth/change-password", json={"current": "admin", "new": "a-long-new-password"}).status_code == 200
    return p


def set_confidence(db, vid, conf="medium"):
    """Test helper: an approved valuation of the given confidence (the fake site gives 'low' after approval)."""
    row = db.get_valuation(vid)
    res = row["result"]
    res["svi"]["triangulation"]["confidence"] = conf
    db.exec("UPDATE studio.valuations SET result=%s WHERE id=%s", (Jsonb(res), vid))


# ================================================================== flag off
@needs_db
def test_flag_off_endpoints_404_and_v3_unchanged(v5_off, studio_env):
    env = studio_env
    u, a = env["client"](), env["client"]()
    siwe_login(u, USER_KEY)
    siwe_login(a, ADMIN_KEY)
    vid = run_valuation(env, u)
    v = u.get(f"/v1/studio/valuations/{vid}").json()
    assert v["svi"]["triangulation"]["version"] == "v3"
    assert [s["key"] for s in v["steps"]] == ["read_site", "profile", "competitors", "market", "svi", "narrative"]
    row = env["db"].get_valuation(vid)
    assert "valuation_inputs" not in row["result"] and row["final"] is None
    tri = row["result"]["svi"]["triangulation"]
    assert not {"params_version", "football_field", "tokenisation", "valuation_class"} & set(tri)
    for path in (f"/v1/studio/valuations/{vid}/tokenisation", f"/v1/studio/valuations/{vid}/projections",
                 "/v1/studio/projection-template"):
        assert u.get(path).status_code == 404
    assert u.post(f"/v1/studio/valuations/{vid}/finalise", json={}).status_code == 404
    assert a.get("/v1/admin/price-requests").status_code == 404
    assert fin.company_create_guard(env["db"], row, None) is None


# ================================================================== full v5 flow
@needs_db
def test_v5_projections_finalise_company(v5_on, studio_env):
    env, db = studio_env, studio_env["db"]
    u, a, o = env["client"](), env["client"](), env["client"]()
    siwe_login(u, USER_KEY)
    siwe_login(a, ADMIN_KEY)
    siwe_login(o, "0x" + "c3" * 32)
    vid = run_valuation(env, u)
    v = u.get(f"/v1/studio/valuations/{vid}").json()
    assert v["status"] == "waiting_approval" and v["svi"]["triangulation"]["version"] == "v5"
    assert "valuation_methods" in [s["key"] for s in v["steps"]]
    assert next(s for s in v["steps"] if s["key"] == "valuation_methods")["status"] == "done"
    assert db.get_valuation(vid)["result"]["valuation_inputs"]["industry"]  # keyword fallback (fake LLM)

    # ---- template + upload (multipart xlsx) + permissions + checks
    t = u.get("/v1/studio/projection-template?format=xlsx")
    assert t.status_code == 200 and t.content[:2] == b"PK" and "attachment" in t.headers["content-disposition"]
    assert u.get("/v1/studio/projection-template?format=csv&lang=vi").text.startswith("key,label")
    body, ct = multipart(forecast_xlsx(), "forecast.xlsx")
    assert o.post(f"/v1/studio/valuations/{vid}/projections", content=body,
                  headers={"Content-Type": ct}).status_code == 403
    r = u.post(f"/v1/studio/valuations/{vid}/projections", content=body, headers={"Content-Type": ct})
    assert r.status_code == 201, r.text
    pv = r.json()
    assert pv["status"] == "draft" and pv["errors"] == 0 and pv["can_confirm"] and len(pv["sha256"]) == 64
    assert pv["parsed"]["basis"] == "management_projection" and len(pv["parsed"]["years"]) == 7
    # refused files
    bad, ct2 = multipart(forecast_xlsx(), "forecast.xlsm")
    assert u.post(f"/v1/studio/valuations/{vid}/projections", content=bad, headers={"Content-Type": ct2}
                  ).status_code == 415
    big = u.post(f"/v1/studio/valuations/{vid}/projections?filename=x.csv", content=b"a" * (600 * 1024),
                 headers={"Content-Type": "text/csv"})
    assert big.status_code == 413
    # JSON grid with an identity error: stored as draft, cannot be confirmed
    grid = {"currency": "AUD", "years": [{"year": 2026 + i, "actual": i == 0, "revenue": 1e6, "cogs": 1e5,
                                          "opex": 5e5, "ebitda": 1.0, "d_and_a": 1e4, "capex": 1e4}
                                         for i in range(4)]}
    gj = u.post(f"/v1/studio/valuations/{vid}/projections", json=grid)
    assert gj.status_code == 201 and gj.json()["errors"] >= 1 and not gj.json()["can_confirm"]
    assert u.post(f"/v1/studio/valuations/{vid}/projections/{gj.json()['id']}/confirm",
                  json={"attest": True}).status_code == 409
    assert u.delete(f"/v1/studio/valuations/{vid}/projections/{gj.json()['id']}").json() == {"ok": True}
    lst = u.get(f"/v1/studio/valuations/{vid}/projections").json()
    assert lst["latest"]["id"] == pv["id"] and lst["label"].startswith("Based on management projections")
    exp = u.get(f"/v1/studio/valuations/{vid}/projections/{pv['id']}/export.csv")
    assert exp.status_code == 200 and exp.text.startswith("key,label")

    # ---- confirm: attestation required, then deterministic re-value with projection methods
    assert u.post(f"/v1/studio/valuations/{vid}/projections/{pv['id']}/confirm", json={}).status_code == 422
    c = u.post(f"/v1/studio/valuations/{vid}/projections/{pv['id']}/confirm", json={"attest": True})
    assert c.status_code == 200, c.text
    cj = c.json()
    assert cj["projection"]["status"] == "confirmed" and cj["valuation_status"] == "waiting_approval"
    tri = db.get_valuation(vid)["result"]["svi"]["triangulation"]
    methods = {m["method"] for m in tri["methods"]}
    assert methods & {"vc_method", "first_chicago"} and tri["projections"]["sha256"] == pv["sha256"]
    audit = [x["action"] for x in db.all("SELECT action FROM studio.audit WHERE target=%s", (vid,))]
    assert {"projection_uploaded", "projection_confirmed"} <= set(audit)

    # ---- approval keeps v5 (stored-result path) and confirms; tokenisation proposal
    tk = u.get(f"/v1/studio/valuations/{vid}/tokenisation").json()
    assert not tk["can_finalise"] and any("approve" in b for b in tk["blockers"])
    assert a.post(f"/v1/studio/valuations/{vid}/decision", json={"approved": True}).status_code == 200
    row = db.get_valuation(vid)
    assert row["status"] == "approved" and row["result"]["svi"]["triangulation"]["version"] == "v5"
    assert row["result"]["valuation_inputs"]["projection"]["id"] == pv["id"]
    # /verify path recomputes the stored v5 report exactly
    rep = json.loads(canonical_json(report_view_from_row(row)))
    rc = verify.recompute(rep)
    assert rc["formula_version"] == "v5" and rc["matches_report"]["mid"] and rc["matches_report"]["low"] \
        and rc["matches_report"]["high"] and rc["methods"]

    # ---- low confidence: owner refused, admin may override with a reason
    set_confidence(db, vid, "low")
    r = u.post(f"/v1/studio/valuations/{vid}/finalise", json={})
    assert r.status_code == 409 and "confidence" in r.json()["detail"]
    assert a.post(f"/v1/studio/valuations/{vid}/finalise", json={"allow_low_confidence": True,
                                                                "override_reason": "short"}).status_code == 409
    set_confidence(db, vid, "medium")

    # ---- owner: within +/-20 % needs a note; beyond needs a reason and a platform admin
    rec = u.get(f"/v1/studio/valuations/{vid}/tokenisation").json()["proposal"]["recommended_price_per_share_aud"]
    assert rec == pytest.approx(0.25, abs=0.002) or rec == pytest.approx(0.10, abs=0.002) or rec > 0
    r = u.post(f"/v1/studio/valuations/{vid}/finalise", json={"price_per_share_aud": round(rec * 1.1, 4)})
    assert r.status_code == 422 and "note" in r.json()["detail"]
    r = u.post(f"/v1/studio/valuations/{vid}/finalise", json={"price_per_share_aud": round(rec * 6, 4),
                                                             "reason": "x" * 30})
    assert r.status_code == 422  # beyond the hard limit (5x)
    r = u.post(f"/v1/studio/valuations/{vid}/finalise", json={"price_per_share_aud": round(rec * 1.5, 4),
                                                             "reason": "too short"})
    assert r.status_code == 422 and "reason" in r.json()["detail"]
    r = u.post(f"/v1/studio/valuations/{vid}/finalise",
               json={"price_per_share_aud": round(rec * 1.5, 4), "reason": "Our last SAFE converted at this price."})
    assert r.status_code == 202 and r.json()["pending_request"]["status"] == "pending"
    rid = r.json()["pending_request"]["id"]
    assert u.post(f"/v1/studio/valuations/{vid}/finalise", json={}).status_code == 409  # one at a time
    assert u.get("/v1/admin/price-requests").status_code == 403
    q = a.get("/v1/admin/price-requests").json()
    assert [x["id"] for x in q] == [rid] and q[0]["url"]
    assert a.post(f"/v1/admin/price-requests/{rid}/reject", json={"reason": "no"}).status_code == 422
    assert a.post(f"/v1/admin/price-requests/{rid}/reject", json={"reason": "not supported by evidence"}
                  ).status_code == 200
    # within the band with a note: finalised at once, recommended and chosen both stored
    r = u.post(f"/v1/studio/valuations/{vid}/finalise", json={"price_per_share_aud": round(rec * 1.1, 4),
                                                             "note": "round number"})
    assert r.status_code == 200, r.text
    f = r.json()["final"]
    assert r.json()["final_state"] == "valid" and f["recommended_price_per_share_aud"] == rec
    assert f["price_per_share_aud"] == round(rec * 1.1, 4) and f["price_deviation_pct"] == pytest.approx(10, abs=0.1)
    assert f["total_shares"] == round(f["pre_money_aud"] / f["price_per_share_aud"])
    assert f["based_on_projections"] and f["report_hash"] == report_hash_from_row(db.get_valuation(vid))
    # no planned raise sent: the confirmed projections' planned raise (A$1M) at the chosen price
    assert f["planned_raise_aud"] == 1_000_000 and f["planned_raise_source"] == "projections"
    assert f["new_shares"] == int(1_000_000 // f["price_per_share_aud"])
    assert f["dilution_pct"] == round(f["new_shares"] * 100 / (f["total_shares"] + f["new_shares"]), 2)
    valid_until = datetime.fromisoformat(f["valid_until"])
    assert timedelta(days=89) < valid_until - datetime.now(UTC) <= timedelta(days=90)
    assert u.post(f"/v1/studio/valuations/{vid}/finalise", json={}).status_code == 409  # already finalised

    # ---- company creation: defaults from the final; a different price / share count is refused
    holders = [{"name": "Founder", "wallet": USER, "pct": 100}]
    r = u.post("/v1/studio/companies", json={"valuation_id": vid, "name": "AgriTrace", "ticker": "AGX",
                                             "share_price_aud": 1.0, "holders": holders})
    assert r.status_code == 422 and "finalised price" in r.json()["detail"]
    r = u.post("/v1/studio/companies", json={"valuation_id": vid, "name": "AgriTrace", "ticker": "AGX",
                                             "total_shares": 12345, "holders": holders})
    assert r.status_code == 422
    # expired final -> 409; restored -> created with the final's values
    fx = dict(f, valid_until=(datetime.now(UTC) - timedelta(days=1)).isoformat())
    db.exec("UPDATE studio.valuations SET final=%s WHERE id=%s", (Jsonb(fx), vid))
    assert u.get(f"/v1/studio/valuations/{vid}/tokenisation").json()["final_state"] == "expired"
    assert u.post("/v1/studio/companies", json={"valuation_id": vid, "name": "AgriTrace", "ticker": "AGX",
                                                "holders": holders}).status_code == 409
    db.exec("UPDATE studio.valuations SET final=%s WHERE id=%s", (Jsonb(f), vid))
    r = u.post("/v1/studio/companies", json={"valuation_id": vid, "name": "AgriTrace", "ticker": "AGX",
                                             "holders": holders})
    assert r.status_code == 201, r.text
    comp = r.json()
    assert float(comp["share_price_aud"]) == pytest.approx(f["price_per_share_aud"])
    assert int(comp["total_shares"]) == f["total_shares"] and float(comp["valuation_aud"]) == f["pre_money_aud"]
    # once a company exists: projections and assumptions are frozen
    assert a.patch(f"/v1/admin/valuations/{vid}/assumptions",
                   json={"changes": {"dcf.g": 0.02}, "reason": "house view on growth"}).status_code == 409

    # ---- offering guards (unit level on the real DB rows)
    c = db.one("SELECT * FROM studio.companies WHERE id=%s", (comp["id"],))
    fin.offering_submit_guard(db, c, f["price_per_share_aud"])  # ok
    with pytest.raises(HTTPException) as e:
        fin.offering_submit_guard(db, c, f["high_aud"] / f["total_shares"] * 1.5)
    assert e.value.status_code == 409 and "top of the finalised range" in e.value.detail
    price, final = fin.offering_default_price(db, c, 1.0, datetime.now(UTC) - timedelta(days=1))
    assert float(price) == f["price_per_share_aud"] and final["valuation_id"] == vid
    assert float(fin.offering_default_price(db, c, 1.0, datetime.now(UTC) + timedelta(days=1))[0]) == 1.0
    extra = fin.pack_valuation_extra(db, c)
    assert extra["final"]["price_per_share_aud"] == f["price_per_share_aud"] and extra["football_field"]
    assert extra["based_on_projections"] and extra["projection_sha256"] == pv["sha256"]
    # KPI revenue far from the revenue used -> revaluation needed
    if f.get("revenue_used_aud"):
        db.exec("INSERT INTO studio.kpi_values (company_id, period_end, metric, value, cadence) VALUES "
                "(%s, '2026-08-31', 'revenue', %s, 'monthly')", (c["id"], f["revenue_used_aud"] / 12 * 2))
        with pytest.raises(HTTPException) as e:
            fin.offering_submit_guard(db, c, f["price_per_share_aud"])
        assert "revaluation needed" in e.value.detail
        db.exec("DELETE FROM studio.kpi_values WHERE company_id=%s", (c["id"],))
    # stage changed since finalising
    db.exec("UPDATE studio.valuations SET final=%s WHERE id=%s", (Jsonb(dict(f, stage="growth")), vid))
    with pytest.raises(HTTPException) as e:
        fin.offering_submit_guard(db, c, f["price_per_share_aud"])
    assert "stage changed" in e.value.detail
    # a newer approved valuation of the same website
    db.exec("UPDATE studio.valuations SET final=%s WHERE id=%s", (Jsonb(f), vid))
    db.exec("INSERT INTO studio.valuations (id, url, requested_by, status, result, updated_at) VALUES "
            "('newer', %s, 'x', 'approved', '{}', now())", (row["url"],))
    with pytest.raises(HTTPException) as e:
        fin.offering_submit_guard(db, c, f["price_per_share_aud"])
    assert "newer approved valuation" in e.value.detail


@needs_db
def test_v5_price_approval_four_eyes_and_assumptions(v5_on, studio_env):
    env, db = studio_env, studio_env["db"]
    a = env["client"]()
    siwe_login(a, ADMIN_KEY)
    p = second_admin(env)
    vid = run_valuation(env, a)  # requested by the ADMIN wallet
    assert a.post(f"/v1/studio/valuations/{vid}/decision", json={"approved": True}).status_code == 200
    set_confidence(db, vid, "medium")
    rec = a.get(f"/v1/studio/valuations/{vid}/tokenisation").json()["proposal"]["recommended_price_per_share_aud"]
    r = a.post(f"/v1/studio/valuations/{vid}/finalise",
               json={"price_per_share_aud": round(rec * 0.5, 4), "reason": "Board resolution: half the price.",
                     "planned_raise_aud": 750_000})
    assert r.status_code == 202
    rid = r.json()["pending_request"]["id"]
    # the requester (even a platform admin) cannot approve their own request
    assert a.post(f"/v1/admin/price-requests/{rid}/approve", json={}).status_code == 403
    r = p.post(f"/v1/admin/price-requests/{rid}/approve", json={"note": "checked the minutes"})
    assert r.status_code == 200, r.text
    f = r.json()["final"]
    assert f["price_request_id"] == rid and f["price_approved_by"] == "admin"
    assert f["price_per_share_aud"] == round(rec * 0.5, 4) and f["price_deviation_pct"] == pytest.approx(-50, abs=0.1)
    # the founder's planned raise survives the four-eyes request
    assert f["planned_raise_aud"] == 750_000 and f["planned_raise_source"] == "founder"
    assert f["new_shares"] == int(750_000 // f["price_per_share_aud"])
    assert f["post_money_aud"] == pytest.approx(f["price_per_share_aud"] * f["total_shares"] + 750_000, abs=0.01)
    assert a.post(f"/v1/studio/valuations/{vid}/finalise", json={"planned_raise_aud": -1}).status_code == 422
    acts = {x["action"] for x in db.all("SELECT action FROM studio.audit WHERE target=%s", (vid,))}
    assert {"valuation_price_requested", "valuation_price_approved", "valuation_finalised"} <= acts

    # ---- assumptions: bounds + reason; finalised -> four-eyes pending, second admin applies, final cleared
    got = a.get(f"/v1/admin/valuations/{vid}/assumptions").json()
    assert "dcf.g" in got["allowed"] and got["final_state"] == "valid"
    assert a.patch(f"/v1/admin/valuations/{vid}/assumptions",
                   json={"changes": {"dcf.g": 0.2}, "reason": "house view on growth"}).status_code == 422
    assert a.patch(f"/v1/admin/valuations/{vid}/assumptions",
                   json={"changes": {"nope": 1}, "reason": "house view on growth"}).status_code == 422
    assert a.patch(f"/v1/admin/valuations/{vid}/assumptions",
                   json={"changes": {"scorecard.base_pre_money": 5e6}, "reason": "short"}).status_code == 422
    r = a.patch(f"/v1/admin/valuations/{vid}/assumptions",
                json={"changes": {"scorecard.base_pre_money": 5e6}, "reason": "AU seed median is lower in 2026"})
    assert r.status_code == 200 and r.json()["pending"] and not r.json()["applied"]
    cid = r.json()["change_ids"][0]
    assert a.post(f"/v1/admin/assumption-changes/{cid}/approve").status_code == 403  # same admin
    r = p.post(f"/v1/admin/assumption-changes/{cid}/approve")
    assert r.status_code == 200, r.text
    row = db.get_valuation(vid)
    assert row["final"] is None and row["status"] == "waiting_approval"
    sc = next(m for m in row["result"]["svi"]["triangulation"]["methods"] if m["method"] == "scorecard")
    assert sc["inputs"]["base_pre_money_aud"] == 5e6
    ch = db.one("SELECT * FROM studio.valuation_assumption_changes WHERE id=%s", (cid,))
    assert ch["status"] == "applied" and ch["approved_by"] == "admin" and ch["report_hash_after"]

    # not finalised: applied at once, audit row per change, stage override recorded as human
    r = a.patch(f"/v1/admin/valuations/{vid}/assumptions",
                json={"changes": {"stage": "pre-seed", "confirm_startup_factors": True},
                      "reason": "no product revenue yet, website overstates"})
    assert r.status_code == 200 and r.json()["applied"], r.text
    tri = db.get_valuation(vid)["result"]["svi"]["triangulation"]
    assert tri["stage"]["stage"] == "pre-seed" and tri["stage"]["basis"] == "human"
    assert tri["tokenisation"]["default_price_for_stage_aud"] == 0.10
    n = db.one("SELECT count(*) AS n FROM studio.valuation_assumption_changes WHERE valuation_id=%s AND "
               "status='applied'", (vid,))["n"]
    assert n == 3
    # re-value an old report with v5 (admin, deterministic)
    assert a.post(f"/v1/admin/valuations/{vid}/rerun-v5", json={"reason": "short"}).status_code == 422
    r = a.post(f"/v1/admin/valuations/{vid}/rerun-v5", json={"reason": "re-run with the v5 standard methods"})
    assert r.status_code == 200 and r.json()["valuation"]["version"] == "v5"


@needs_db
def test_require_final_flag_blocks_company_creation(v5_on, studio_env, monkeypatch):
    env = studio_env
    a = env["client"]()
    siwe_login(a, ADMIN_KEY)
    vid = run_valuation(env, a)
    assert a.post(f"/v1/studio/valuations/{vid}/decision", json={"approved": True}).status_code == 200
    monkeypatch.setenv("VALUATION_V5_REQUIRE_FINAL", "1")
    from test_studio import ADMIN

    r = a.post("/v1/studio/companies", json={"valuation_id": vid, "name": "X Co", "ticker": "XCO",
                                             "holders": [{"name": "A", "wallet": ADMIN, "pct": 100}]})
    assert r.status_code == 409 and "finalise" in r.json()["detail"]
    monkeypatch.setenv("VALUATION_V5_REQUIRE_FINAL", "0")
    r = a.post("/v1/studio/companies", json={"valuation_id": vid, "name": "X Co", "ticker": "XCO",
                                             "holders": [{"name": "A", "wallet": ADMIN, "pct": 100}]})
    assert r.status_code == 201  # rollout default: no final -> previous behaviour (A$1 default price)
    assert float(r.json()["share_price_aud"]) == 1.0
