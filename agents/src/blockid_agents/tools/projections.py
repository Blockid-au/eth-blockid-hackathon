"""Management projections: template, parsing and validation (docs/PLAN-VALUATION-V5.md §6). No LLM, no formulas.

This is the "package reader" tier: an uploaded spreadsheet is untrusted. It is read here with fixed rules only and
never shown to a model (the Anthropic `valuation-reviewer` isolation idea). Guards, in order:
  1. size cap (512 KB) before anything else; extension + magic bytes (`PK\\x03\\x04` for .xlsx, text for .csv);
     .xlsm / .xls / macro-enabled content types / vbaProject.bin rejected
  2. zip-bomb guard: <= 200 entries, <= 20 MB uncompressed in total (from the zip directory, before openpyxl)
  3. openpyxl read-only with defusedxml installed (XML entity attacks refused; we fail closed without it),
     `data_only=True` (cached values; formulas are never evaluated); a formula cell without a cached value is an
     error ("open and save the file in Excel first"); first 3 sheets, 200 rows x 20 columns at most
  4. text fields: control characters stripped, length-capped; numbers: finite, plain (1,200 / (300) / 1.2e6)
  5. CSV export of stored projections escapes cells starting with = + - @ (formula injection)

Layout (template_version 1), sheet "Projections" (CSV: the same table; company fields as rows `company.<key>`):
    key | label | FY1 | FY2 | ... (up to 3 actual + 5 projected years)
    year, actual (A/P), revenue, cogs, opex, ebitda (optional check), d_and_a, tax (optional), capex,
    nwc or change_nwc (optional), headcount (optional), customers (optional)
Sheet "Company": key | label | value — currency, fiscal_year_end, cash, debt, shares_fd, planned_raise, audited (Y/N),
prepared_by, basis_notes.
"""
from __future__ import annotations

import csv
import io
import math
import re
import zipfile
from typing import Any

from ..config import FX_TO_AUD, FX_TO_AUD_AS_OF
from ..schemas import Check, ProjectionInput
from . import valuation_methods as vm

TEMPLATE_VERSION = 1
YEAR_ROWS = ("year", "actual", "revenue", "cogs", "opex", "ebitda", "d_and_a", "tax", "capex", "nwc", "change_nwc",
             "headcount", "customers")
REQUIRED_ROWS = ("year", "actual", "revenue", "cogs", "opex", "d_and_a", "capex")
COMPANY_ROWS = ("currency", "fiscal_year_end", "cash", "debt", "shares_fd", "planned_raise", "audited", "prepared_by",
                "basis_notes")
LABELS = {
    "en": {"year": "Fiscal year ending (YYYY)", "actual": "Actual (A) or Projection (P)", "revenue": "Revenue",
           "cogs": "Cost of sales", "opex": "Operating expenses (excluding depreciation & amortisation)",
           "ebitda": "EBITDA (optional check: revenue - cost of sales - operating expenses)",
           "d_and_a": "Depreciation & amortisation", "tax": "Income tax paid (optional)",
           "capex": "Capital expenditure", "nwc": "Net working capital balance (optional)",
           "change_nwc": "Change in net working capital (optional)", "headcount": "Employees, full-time equivalent "
           "(optional)", "customers": "Paying customers (optional)",
           "currency": "Currency (AUD, USD, VND, ...)", "fiscal_year_end": "Fiscal year end (MM-DD)",
           "cash": "Cash today", "debt": "Debt today", "shares_fd": "Shares on issue, fully diluted (optional)",
           "planned_raise": "Planned raise (optional)", "audited": "Actual years audited? (Y/N)",
           "prepared_by": "Prepared by", "basis_notes": "Basis of preparation (notes, not used in the maths)"},
    "vi": {"year": "Năm tài chính kết thúc (YYYY)", "actual": "Thực tế (A) hoặc Dự báo (P)", "revenue": "Doanh thu",
           "cogs": "Giá vốn hàng bán", "opex": "Chi phí hoạt động (không gồm khấu hao)",
           "ebitda": "EBITDA (tùy chọn, để đối chiếu)", "d_and_a": "Khấu hao", "tax": "Thuế TNDN đã nộp (tùy chọn)",
           "capex": "Chi đầu tư tài sản", "nwc": "Vốn lưu động ròng (tùy chọn)",
           "change_nwc": "Thay đổi vốn lưu động ròng (tùy chọn)", "headcount": "Số nhân viên (tùy chọn)",
           "customers": "Khách hàng trả tiền (tùy chọn)", "currency": "Tiền tệ (AUD, USD, VND, ...)",
           "fiscal_year_end": "Ngày kết thúc năm tài chính (MM-DD)", "cash": "Tiền mặt hiện có",
           "debt": "Nợ vay hiện có", "shares_fd": "Số cổ phần (pha loãng hoàn toàn, tùy chọn)",
           "planned_raise": "Số vốn dự kiến gọi (tùy chọn)", "audited": "Số liệu thực tế đã kiểm toán? (Y/N)",
           "prepared_by": "Người lập", "basis_notes": "Cơ sở lập (ghi chú, không dùng để tính)"},
}
GUIDE = {
    "en": ["How to fill this template (BlockID valuation, template version 1)",
           "1. One column per fiscal year: up to 3 ACTUAL years (A) first, then 3 to 5 PROJECTED years (P).",
           "2. Amounts in the currency on the Company sheet, whole units (no thousands or millions).",
           "3. Type numbers, not formulas. If you use formulas, open and save the file in Excel before uploading.",
           "4. These are your management's projections: BlockID checks them against simple limits but does not "
           "verify them. The valuation will say 'Based on management projections (unaudited, not verified by "
           "BlockID)'.", "5. Upload .xlsx or .csv, at most 512 KB. Macro files (.xlsm) are refused."],
    "vi": ["Cách điền mẫu này (định giá BlockID, phiên bản mẫu 1)",
           "1. Mỗi cột là một năm tài chính: tối đa 3 năm THỰC TẾ (A) trước, sau đó 3 đến 5 năm DỰ BÁO (P).",
           "2. Số tiền theo đơn vị tiền tệ ở trang Company, ghi đủ số (không viết tắt nghìn/triệu).",
           "3. Nhập số, không nhập công thức. Nếu dùng công thức, hãy mở và lưu tệp trong Excel trước khi tải lên.",
           "4. Đây là dự báo của ban điều hành: BlockID kiểm tra theo giới hạn đơn giản nhưng không xác minh. "
           "Định giá sẽ ghi 'Dựa trên dự báo của ban điều hành (chưa kiểm toán, BlockID chưa xác minh)'.",
           "5. Tải lên tệp .xlsx hoặc .csv, tối đa 512 KB. Tệp có macro (.xlsm) bị từ chối."],
}
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_NUM = re.compile(r"^\(?-?[0-9][0-9,  ]*(\.[0-9]+)?([eE][+-]?[0-9]+)?\)?$")


class ProjectionError(ValueError):
    """Upload refused before validation (size, type, bomb, unreadable). Message is plain words for the user."""


# ------------------------------------------------------------------ helpers
def clean_text(v: Any, limit: int = 500) -> str:
    return _CTRL.sub("", str(v if v is not None else "")).strip()[:limit]


def to_number(v: Any) -> float | None:
    """int/float/plain numeric string -> float; '' / None -> None; anything else raises ValueError."""
    if v is None:
        return None
    if isinstance(v, bool):
        raise ValueError("not a number")
    if isinstance(v, int | float):
        f = float(v)
        if not math.isfinite(f):
            raise ValueError("not a finite number")
        return f
    s = clean_text(v, 64).replace(" ", " ")
    if s == "" or s in ("-", "—"):
        return None
    if not _NUM.match(s):
        raise ValueError(f"'{s[:20]}' is not a plain number")
    neg = s.startswith("(") and s.endswith(")")
    f = float(s.strip("()").replace(",", "").replace(" ", "").replace(" ", ""))
    if not math.isfinite(f):
        raise ValueError("not a finite number")
    return -f if neg else f


def escape_csv_cell(v: Any) -> str:
    s = "" if v is None else str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


# ------------------------------------------------------------------ template
def template_rows(lang: str = "en") -> tuple[list[list[Any]], list[list[Any]]]:
    lab = LABELS.get(lang, LABELS["en"])
    years = [2024, 2025, 2026, 2027, 2028, 2029, 2030]
    kinds = ["A", "A", "P", "P", "P", "P", "P"]
    proj = [["key", "label", *[f"FY{i + 1}" for i in range(len(years))]]]
    for k in YEAR_ROWS:
        if k == "year":
            proj.append([k, lab[k], *years])
        elif k == "actual":
            proj.append([k, lab[k], *kinds])
        else:
            proj.append([k, lab[k], *([None] * len(years))])
    comp = [["key", "label", "value"], ["currency", lab["currency"], "AUD"],
            ["fiscal_year_end", lab["fiscal_year_end"], "06-30"]]
    comp += [[k, lab[k], "N" if k == "audited" else None] for k in COMPANY_ROWS if k not in ("currency",
                                                                                            "fiscal_year_end")]
    return proj, comp


def template_xlsx(lang: str = "en") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    proj, comp = template_rows(lang)
    wb = Workbook()
    ws = wb.active
    ws.title = "Projections"
    for r in proj:
        ws.append(r)
    blue = Font(color="1F4E9E")
    for row in ws.iter_rows(min_row=2, min_col=3):
        for c in row:
            c.font = blue  # input cells (skill convention: inputs blue)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF7")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 58
    wc = wb.create_sheet("Company")
    for r in comp:
        wc.append(r)
    wc.column_dimensions["B"].width = 50
    wg = wb.create_sheet("Guide")
    for line in GUIDE.get(lang, GUIDE["en"]):
        wg.append([line])
    wg.column_dimensions["A"].width = 120
    wb.properties.title = f"BlockID projections template v{TEMPLATE_VERSION}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def template_csv(lang: str = "en") -> bytes:
    proj, comp = template_rows(lang)
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in proj:
        w.writerow(["" if x is None else x for x in r])
    for r in comp[1:]:
        w.writerow([f"company.{r[0]}", r[1], "" if r[2] is None else r[2]])
    return buf.getvalue().encode("utf-8")


# ------------------------------------------------------------------ reading
def sniff(data: bytes, filename: str, p) -> str:
    """'xlsx' | 'csv' or ProjectionError. Size first, then extension, then content."""
    lim = int(p["projection"]["max_bytes"])
    if len(data) > lim:
        raise ProjectionError(f"the file is larger than {lim // 1024} KB")
    if not data:
        raise ProjectionError("the file is empty")
    name = (filename or "").lower().strip()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    if ext in ("xlsm", "xls", "xlsb", "xltm", "ods"):
        raise ProjectionError("only .xlsx or .csv files are accepted (macro-enabled and old Excel files are not)")
    if data[:4] == b"PK\x03\x04":
        if ext not in ("xlsx", ""):
            raise ProjectionError("the file content is a spreadsheet but its name does not end in .xlsx")
        return "xlsx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ProjectionError("old Excel (.xls) files are not accepted: save as .xlsx")
    if ext not in ("csv", "txt", ""):
        raise ProjectionError("only .xlsx or .csv files are accepted")
    if b"\x00" in data:
        raise ProjectionError("the file is not a text CSV")
    return "csv"


def zip_guard(data: bytes, p) -> None:
    pp = p["projection"]
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except (zipfile.BadZipFile, ValueError, OSError) as e:
        raise ProjectionError("the .xlsx file is damaged or not a real Excel file") from e
    if len(infos) > pp["max_zip_entries"]:
        raise ProjectionError("the .xlsx file has too many parts")
    total = sum(max(i.file_size, 0) for i in infos)
    if total > pp["max_uncompressed_bytes"]:
        raise ProjectionError("the .xlsx file expands to more than 20 MB: refused")
    names = {i.filename for i in infos}
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise ProjectionError("files with macros are not accepted")
    if "xl/workbook.xml" not in names:
        raise ProjectionError("the file is not an Excel workbook (.xlsx)")
    try:
        ct = zf.read("[Content_Types].xml")[:200_000].decode("utf-8", "replace").lower()
    except KeyError:
        raise ProjectionError("the .xlsx file is damaged") from None
    if "macroenabled" in ct or "vbaproject" in ct:
        raise ProjectionError("files with macros are not accepted")


def _xlsx_tables(data: bytes, p) -> dict[str, list[list[Any]]]:
    import openpyxl
    from openpyxl.xml import DEFUSEDXML

    if not DEFUSEDXML:  # fail closed: never parse untrusted XML without defusedxml
        raise ProjectionError("spreadsheet uploads are unavailable on this server (XML guard missing); use CSV")
    pp = p["projection"]
    zip_guard(data, p)
    try:
        wb_v = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        wb_f = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as e:
        raise ProjectionError("the .xlsx file could not be read; save it again as an Excel workbook") from e
    out: dict[str, list[list[Any]]] = {}
    try:
        for idx, name in enumerate(wb_v.sheetnames[: pp["max_sheets"]]):
            ws_v, ws_f = wb_v[name], wb_f[name]
            rows_v = list(ws_v.iter_rows(max_row=pp["max_rows"], max_col=pp["max_cols"], values_only=True))
            rows_f = list(ws_f.iter_rows(max_row=pp["max_rows"], max_col=pp["max_cols"], values_only=True))
            table = []
            for i, rv in enumerate(rows_v):
                rf = rows_f[i] if i < len(rows_f) else ()
                row = []
                for j, v in enumerate(rv):
                    f = rf[j] if j < len(rf) else None
                    if v is None and isinstance(f, str) and f.startswith("="):
                        raise ProjectionError(f"sheet '{clean_text(name, 40)}' cell {_cell(i, j)} has a formula with "
                                              "no saved result: open and save the file in Excel, then upload again")
                    row.append(v)
                table.append(row)
            out[clean_text(name, 40).lower() or f"sheet{idx}"] = table
    finally:
        wb_v.close()
        wb_f.close()
    return out


def _cell(i: int, j: int) -> str:
    col = ""
    j += 1
    while j:
        j, r = divmod(j - 1, 26)
        col = chr(65 + r) + col
    return f"{col}{i + 1}"


def _csv_tables(data: bytes, p) -> dict[str, list[list[Any]]]:
    pp = p["projection"]
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1252")
        except UnicodeDecodeError as e:
            raise ProjectionError("the CSV file is not UTF-8 text") from e
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[Any]] = []
    for i, r in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= pp["max_rows"]:
            break
        rows.append([x for x in r[: pp["max_cols"]]])
    proj = [r for r in rows if not (r and str(r[0]).strip().lower().startswith("company."))]
    comp = [["key", "label", "value"]] + [[str(r[0]).strip()[8:], r[1] if len(r) > 1 else "",
                                           r[2] if len(r) > 2 else ""]
                                          for r in rows if r and str(r[0]).strip().lower().startswith("company.")]
    return {"projections": proj, "company": comp}


def read_tables(data: bytes, filename: str, p) -> tuple[str, dict[str, list[list[Any]]]]:
    kind = sniff(data, filename, p)
    return kind, (_xlsx_tables(data, p) if kind == "xlsx" else _csv_tables(data, p))


def _key(v: Any) -> str:
    return re.sub(r"[^a-z_.]", "", clean_text(v, 40).lower().replace(" ", "_"))


def to_input(tables: dict[str, list[list[Any]]], p) -> tuple[dict, list[Check]]:
    """Tables -> ProjectionInput-shaped dict + structure errors (never raises for content problems)."""
    errs: list[Check] = []
    proj = tables.get("projections") or next(iter(tables.values()), [])
    by_key: dict[str, list[Any]] = {}
    for r in proj:
        if not r:
            continue
        k = _key(r[0])
        if k in YEAR_ROWS and k not in by_key:
            by_key[k] = list(r[2:])
    for k in REQUIRED_ROWS:
        if k not in by_key:
            errs.append(Check(code="structure", severity="error", row=k,
                              message=f"Row '{k}' is missing (column A must hold the row key)."))
    if errs:
        return {}, errs
    ncols = max((i + 1 for i, v in enumerate(by_key["year"]) if v not in (None, "")), default=0)
    years: list[dict] = []
    for j in range(ncols):
        y: dict[str, Any] = {}
        try:
            yr = to_number(by_key["year"][j])
        except ValueError:
            yr = None
        if yr is None or yr != int(yr):
            errs.append(Check(code="structure", severity="error", row="year",
                              message=f"Column {_cell(0, j + 2)[:-1]}: the year must be a 4-digit year."))
            continue
        y["year"] = int(yr)
        a = clean_text(by_key["actual"][j] if j < len(by_key["actual"]) else "", 10).upper()
        if a not in ("A", "P", "ACTUAL", "PROJECTION", "F", "FORECAST"):
            errs.append(Check(code="structure", severity="error", year=y["year"], row="actual",
                              message=f"{y['year']}: mark the year A (actual) or P (projection)."))
            continue
        y["actual"] = a in ("A", "ACTUAL")
        empty = True
        for k in YEAR_ROWS[2:]:
            vals = by_key.get(k)
            raw = vals[j] if vals is not None and j < len(vals) else None
            try:
                v = to_number(raw)
            except ValueError as e:
                errs.append(Check(code="structure", severity="error", year=y["year"], row=k,
                                  message=f"{y['year']} '{k}': {e}."))
                v = None
            if v is not None:
                empty = False
            if v is None and k in REQUIRED_ROWS:
                errs.append(Check(code="structure", severity="error", year=y["year"], row=k,
                                  message=f"Row '{k}' is missing for {y['year']}."))
            y[k] = v
        if empty:  # a blank column (e.g. an unused template year) is simply ignored
            errs = [e for e in errs if e.year != y["year"]]
            continue
        years.append(y)
    if not years and not errs:
        errs.append(Check(code="structure", severity="error", row="revenue",
                          message="No year has figures yet: fill in the revenue, costs and capex rows."))
    company: dict[str, Any] = {}
    for r in tables.get("company") or []:
        if not r or len(r) < 3:
            continue
        k = _key(r[0])
        if k in COMPANY_ROWS:
            company[k] = r[2]
    out: dict[str, Any] = {"currency": clean_text(company.get("currency") or "AUD", 8).upper(),
                           "fiscal_year_end": clean_text(company.get("fiscal_year_end") or "06-30", 10),
                           "audited": clean_text(company.get("audited") or "N", 5).upper() in ("Y", "YES", "TRUE"),
                           "prepared_by": clean_text(company.get("prepared_by"), 200),
                           "basis_notes": clean_text(company.get("basis_notes"), int(p["projection"]["text_max"])),
                           "years": years}
    for k in ("cash", "debt", "planned_raise", "shares_fd"):
        try:
            v = to_number(company.get(k))
        except ValueError as e:
            errs.append(Check(code="structure", severity="error", row=k, message=f"Company '{k}': {e}."))
            v = None
        if v is not None:
            if v < 0:
                errs.append(Check(code="structure", severity="error", row=k, message=f"Company '{k}' cannot be negative."))
            out[k] = int(v) if k == "shares_fd" else v
    if out.get("shares_fd") is not None and out["shares_fd"] < 1:
        out.pop("shares_fd")
    return out, errs


# ------------------------------------------------------------------ validation
def validate(inp: ProjectionInput, *, cls: str, industry_row, revenue_ref_aud: float | None, p) -> tuple[list[Check],
                                                                                                           list[dict]]:
    """Checks (plan §6.4) + the years actually used (projected years after caps, in the upload currency)."""
    pp = p["projection"]
    checks: list[Check] = []
    ys = sorted((y.model_dump() for y in inp.years), key=lambda y: y["year"])
    actual = [y for y in ys if y["actual"]]
    projected = [y for y in ys if not y["actual"]]
    years = [y["year"] for y in ys]
    if len(set(years)) != len(years) or any(b - a != 1 for a, b in zip(years, years[1:])):
        checks.append(Check(code="structure", severity="error", message="Years must be consecutive, one column each."))
    if actual and projected and max(y["year"] for y in actual) > min(y["year"] for y in projected):
        checks.append(Check(code="structure", severity="error", message="Actual years must come before projected years."))
    if not pp["min_projected_years"] <= len(projected) <= pp["max_projected_years"]:
        checks.append(Check(code="structure", severity="error",
                            message=f"Give {pp['min_projected_years']} to {pp['max_projected_years']} projected years "
                                    f"(found {len(projected)})."))
    if len(actual) > pp["max_actual_years"]:
        checks.append(Check(code="structure", severity="error", message="Give at most 3 actual years."))
    if inp.currency.upper() not in FX_TO_AUD:
        checks.append(Check(code="currency", severity="error", message=f"Currency {inp.currency} isn't supported."))
    tol = pp["identity_tolerance"]
    for y in ys:
        calc = y["revenue"] - y["cogs"] - y["opex"]
        if y.get("ebitda") is not None and abs(y["ebitda"] - calc) > tol * max(abs(calc), abs(y["ebitda"]), 1.0):
            checks.append(Check(code="identity", severity="error", year=y["year"], row="ebitda",
                                message=f"EBITDA {y['year']} doesn't add up: {y['ebitda']:,.0f} given, {calc:,.0f} "
                                        "from revenue - cost of sales - operating expenses."))
    if any(c.severity == "error" for c in checks):
        return checks, []

    used = [dict(y) for y in projected]
    # growth caps (stage), each later year's cap = previous cap x decay
    cap = float(pp["growth_cap_y1"].get(cls, 1.5))
    prev = actual[-1]["revenue"] if actual else None
    if actual and projected and actual[-1]["revenue"] > 0:
        jump = pp["jump_from_actual_sme"] if cls in ("profitable_sme", "listed") else pp["jump_from_actual"]
        if projected[0]["revenue"] > jump * actual[-1]["revenue"]:
            checks.append(Check(code="jump_from_actual", severity="warning", year=projected[0]["year"], row="revenue",
                                message=f"Revenue in {projected[0]['year']} is more than {jump:g}x the last actual "
                                        "year."))
    for i, y in enumerate(used):
        if prev is not None and prev > 0:
            g = y["revenue"] / prev - 1
            if g > cap:
                new_rev = prev * (1 + cap)
                ratio = new_rev / y["revenue"]
                checks.append(Check(code="growth_cap", severity="warning", year=y["year"], row="revenue",
                                    used_value=round(new_rev, 2),
                                    message=f"Revenue grows {g:.0%} in {y['year']}; we used {cap:.0%} (limit for "
                                            f"this stage)."))
                for k in ("revenue", "cogs", "opex", "d_and_a", "capex", "tax", "nwc", "change_nwc"):
                    if y.get(k) is not None:
                        y[k] = y[k] * ratio
        prev = y["revenue"]
        cap = cap * float(pp["growth_cap_decay"])
    # margins
    band = industry_row["gross_margin_band"]
    for y in used:
        if y["revenue"] > 0:
            gm = (y["revenue"] - y["cogs"]) / y["revenue"]
            if not band[0] <= gm <= band[1]:
                checks.append(Check(code="gross_margin", severity="warning", year=y["year"], row="cogs",
                                    message=f"Gross margin {gm:.0%} in {y['year']} is outside the usual "
                                            f"{band[0]:.0%}-{band[1]:.0%} for {industry_row['label']}."))
                break
    m_cap = float(industry_row["ebitda_margin_p90"]) + float(pp["margin_cap_pp_over_p90"])
    for y in used:
        if y["revenue"] > 0:
            margin = (y["revenue"] - y["cogs"] - y["opex"]) / y["revenue"]
            if margin > m_cap:
                target_opex = y["revenue"] * (1 - m_cap) - y["cogs"]
                checks.append(Check(code="margin_cap", severity="warning", year=y["year"], row="opex",
                                    used_value=round(max(target_opex, 0.0), 2),
                                    message=f"EBITDA margin {margin:.0%} in {y['year']} is above what listed peers "
                                            f"reach ({industry_row['ebitda_margin_p90']:.0%}); we used {m_cap:.0%}."))
                y["opex"] = max(target_opex, 0.0)
    # revenue per employee
    norm = float(pp["revenue_per_fte_norm_aud"]) / float(FX_TO_AUD.get(inp.currency.upper(), 1.0))
    for y in used:
        if y.get("headcount") and y["headcount"] > 0 and y["revenue"] / y["headcount"] > pp["revenue_per_fte_max_ratio"] * norm:
            checks.append(Check(code="revenue_per_fte", severity="warning", year=y["year"], row="headcount",
                                message=f"Revenue per employee in {y['year']} is more than 3x the usual level."))
            break
    da = sum(y["d_and_a"] for y in used)
    if da > 0 and sum(y["capex"] for y in used) < pp["capex_vs_da_min"] * da:
        checks.append(Check(code="capex_vs_da", severity="warning", row="capex",
                            message="Capital expenditure is below half of depreciation over the forecast: the assets "
                                    "would shrink."))
    if any(y.get("tax") is None for y in used):
        checks.append(Check(code="tax_default", severity="info", row="tax",
                            message="Tax not given for some years: company tax rate applied to positive profit."))
    if any(y.get("nwc") is None and y.get("change_nwc") is None for y in used):
        checks.append(Check(code="nwc_default", severity="info", row="nwc",
                            message="Working capital not given: 10 % of the revenue increase assumed."))
    if actual and revenue_ref_aud and revenue_ref_aud > 0:
        last = actual[-1]["revenue"] * float(FX_TO_AUD.get(inp.currency.upper(), 1.0))
        if abs(last / revenue_ref_aud - 1) > pp["actuals_mismatch"]:
            checks.append(Check(code="actuals_mismatch", severity="warning", year=actual[-1]["year"], row="revenue",
                                message=f"Your {actual[-1]['year']} actual revenue (A${last:,.0f}) differs from the "
                                        f"revenue on your valuation (A${revenue_ref_aud:,.0f})."))
    base_rev = actual[-1]["revenue"] if actual else 0.0
    fc = [x["fcff"] for x in vm.fcff_series(used, tax_rate=0.25, base_revenue=base_rev)]
    cum = sum(fc)
    if cum > 0 and fc and fc[-1] >= p["projection"]["hockey_stick_share"] * cum and len(fc) >= 3:
        checks.append(Check(code="hockey_stick", severity="warning", row="revenue",
                            message="Most of the forecast cash flow arrives in the final year (hockey stick)."))
    return checks, [{k: (round(v, 2) if isinstance(v, float) else v) for k, v in y.items()} for y in actual + used]


def parse_upload(data: bytes, filename: str, p) -> tuple[ProjectionInput | None, list[Check], str]:
    """bytes -> (ProjectionInput or None, structure errors, kind). Raises ProjectionError for refused files."""
    kind, tables = read_tables(data, filename, p)
    raw, errs = to_input(tables, p)
    if errs:
        return None, errs, kind
    try:
        return ProjectionInput.model_validate(raw), [], kind
    except Exception as e:  # noqa: BLE001 - pydantic ValidationError -> plain check
        msg = "; ".join(f"{'.'.join(str(x) for x in er.get('loc', ()))}: {er.get('msg')}"
                        for er in getattr(e, "errors", list)()[:5]) or "invalid values"
        return None, [Check(code="structure", severity="error", message=f"Values out of range: {msg}"[:400])], kind


def parsed_record(inp: ProjectionInput, used: list[dict]) -> dict:
    d = inp.model_dump()
    rate = FX_TO_AUD.get(inp.currency.upper())
    d.update(basis="management_projection", fx_rate_to_aud=rate, fx_as_of=FX_TO_AUD_AS_OF,
             net_debt=round(inp.debt - inp.cash, 2), years_used=used, template_version=TEMPLATE_VERSION)
    return d


def export_csv(parsed: dict) -> bytes:
    """Stored projections back to CSV, every cell escaped against formula injection."""
    buf = io.StringIO()
    w = csv.writer(buf)
    ys = parsed.get("years") or []
    w.writerow(["key", "label", *[f"FY{i + 1}" for i in range(len(ys))]])
    for k in YEAR_ROWS:
        vals = [("A" if y.get("actual") else "P") if k == "actual" else y.get(k) for y in ys]
        w.writerow([k, LABELS["en"][k] if k in LABELS["en"] else k, *[escape_csv_cell(v) for v in vals]])
    for k in COMPANY_ROWS:
        w.writerow([f"company.{k}", LABELS["en"][k], escape_csv_cell(parsed.get(k))])
    return buf.getvalue().encode("utf-8")
